import base64
from pathlib import Path
from uuid import uuid4

from docx import Document

from ai_workshop.labs.rag.ocr.contracts import (
    PP_STRUCTURE_V3_RUNTIME_ROLES,
    OcrProfileSpec,
    OcrRequest,
    OcrResult,
    OcrTextUnit,
)
from ai_workshop.labs.rag.parsing.contracts import ParseRequest
from ai_workshop.labs.rag.parsing.docx import DOCX_MEDIA_TYPE, DocxStructureParser
from tests.unit.labs.rag.parsing.test_service import FakeWorkspace

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


class FakeOcrRuntime:
    def __init__(self) -> None:
        self.requests: list[OcrRequest] = []

    def recognize(self, request: OcrRequest, profile: OcrProfileSpec) -> OcrResult:
        self.requests.append(request)
        assert profile.pipeline_name == "PP-StructureV3"
        return OcrResult(
            text_units=(OcrTextUnit("운용 한도 7%", 0.97, (0.1, 0.2, 0.8, 0.5), True),),
            table_cells=(),
        )


def _profile(tmp_path: Path) -> OcrProfileSpec:
    artifact_directories = {
        role: tmp_path / role for role in PP_STRUCTURE_V3_RUNTIME_ROLES
    }
    for directory in artifact_directories.values():
        directory.mkdir()
    return OcrProfileSpec.create(
        pipeline_name="PP-StructureV3",
        pipeline_version="3.7.0",
        model_names={
            role: f"model-{role}" for role in PP_STRUCTURE_V3_RUNTIME_ROLES
        },
        languages=("ko", "en"),
        confidence_threshold=0.8,
        artifact_directories=artifact_directories,
    )


def _document(tmp_path: Path, *, repeated_image: bool = False) -> Path:
    image_path = tmp_path / "source.png"
    image_path.write_bytes(_PNG)
    document = Document()
    document.add_heading("운용 기준", level=1)
    document.add_paragraph("일반 설명")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "표 데이터"
    document.add_picture(str(image_path))
    if repeated_image:
        document.add_picture(str(image_path))
    document.add_paragraph("마지막 문장")
    path = tmp_path / "sample.docx"
    document.save(path)
    return path


def test_docx_preserves_body_order_and_ocr_image_provenance(tmp_path: Path) -> None:
    runtime = FakeOcrRuntime()
    path = _document(tmp_path)

    parsed = DocxStructureParser(
        ocr_runtime=runtime,
        ocr_profile=_profile(tmp_path),
    ).parse(
        ParseRequest(path, DOCX_MEDIA_TYPE, path.name, uuid4(),
                     FakeWorkspace(path.parent), lambda: None)
    )

    assert [item.kind for item in parsed.elements] == [
        "heading",
        "paragraph",
        "table_cell",
        "ocr_text",
        "paragraph",
    ]
    ocr = parsed.elements[3]
    assert ocr.text == "운용 한도 7%"
    assert ocr.section_path == ("운용 기준",)
    assert ocr.location.source_part == "word/media/image1.png"
    assert ocr.location.image_sha256 == runtime.requests[0].image_sha256
    assert ocr.location.bbox == (0.1, 0.2, 0.8, 0.5)


def test_docx_deduplicates_equal_image_ocr_within_asset_version(tmp_path: Path) -> None:
    runtime = FakeOcrRuntime()
    path = _document(tmp_path, repeated_image=True)

    parsed = DocxStructureParser(
        ocr_runtime=runtime,
        ocr_profile=_profile(tmp_path),
    ).parse(
        ParseRequest(path, DOCX_MEDIA_TYPE, path.name, uuid4(),
                     FakeWorkspace(path.parent), lambda: None)
    )

    assert len(runtime.requests) == 1
    assert [element.text for element in parsed.elements].count("운용 한도 7%") == 2


def test_docx_ocr_requires_explicit_workspace(tmp_path: Path) -> None:
    import pytest

    from ai_workshop.labs.rag.parsing.contracts import ParsingError

    runtime = FakeOcrRuntime()
    path = _document(tmp_path)
    with pytest.raises(ParsingError, match="Tracked temporary workspace required"):
        DocxStructureParser(ocr_runtime=runtime, ocr_profile=_profile(tmp_path)).parse(
            ParseRequest(path, DOCX_MEDIA_TYPE, path.name, uuid4())
        )
    assert runtime.requests == []


def test_docx_marks_opaque_writer_before_runtime_and_allocates_once(tmp_path: Path) -> None:
    started = []
    workspace = FakeWorkspace(tmp_path)

    class CheckedRuntime(FakeOcrRuntime):
        def recognize(self, request, profile):
            assert started == [True]
            assert request.image_path.parent == workspace.root
            assert request.image_path.read_bytes() == _PNG
            return super().recognize(request, profile)

    runtime = CheckedRuntime()
    path = _document(tmp_path, repeated_image=True)
    DocxStructureParser(ocr_runtime=runtime, ocr_profile=_profile(tmp_path)).parse(
        ParseRequest(path, DOCX_MEDIA_TYPE, path.name, uuid4(), workspace,
                     lambda: started.append(True))
    )
    assert len(workspace.files) == 1
    assert len(runtime.requests) == 1
    assert workspace.files[0].exists()
