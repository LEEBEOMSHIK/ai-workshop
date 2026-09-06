import base64
from collections.abc import AsyncIterator
from hashlib import sha256
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from docx import Document

from ai_workshop.labs.rag.chunking.contracts import ChunkingConfig
from ai_workshop.labs.rag.chunking.service import StructuralChunker
from ai_workshop.labs.rag.documents.domain import SourceKind
from ai_workshop.labs.rag.ingestion.serialization import serialize_parsed_document
from ai_workshop.labs.rag.ocr.contracts import (
    OcrProfileSpec,
    OcrRequest,
    OcrResult,
    OcrTextUnit,
)
from ai_workshop.labs.rag.parsing.contracts import ParseRequest
from ai_workshop.labs.rag.parsing.docx import DOCX_MEDIA_TYPE, DocxStructureParser
from ai_workshop.labs.rag.search.viewer import ViewerResource, ViewerService
from ai_workshop.platform.assets.storage import StoredObject
from ai_workshop.shared.errors import AppError

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


class _FakeOcrRuntime:
    def recognize(self, request: OcrRequest, profile: OcrProfileSpec) -> OcrResult:
        assert request.source_part == "word/media/image1.png"
        assert profile.pipeline_name == "PP-StructureV3"
        return OcrResult(
            text_units=(
                OcrTextUnit("운용 한도는 순자산의 7%입니다.", 0.98, (0.1, 0.2, 0.9, 0.5), True),
            ),
            table_cells=(),
        )


class _CharacterCounter:
    def count(self, text: str) -> int:
        return len(text)


class _ViewerRepository:
    def __init__(self, actor_id: UUID, resource: ViewerResource) -> None:
        self.actor_id = actor_id
        self.resource = resource

    async def resolve(
        self, *, actor_id: UUID, asset_version_id: UUID, projection_id: UUID
    ) -> ViewerResource | None:
        if (
            actor_id != self.actor_id
            or asset_version_id != self.resource.asset_version_id
            or projection_id != self.resource.projection_id
        ):
            return None
        return self.resource


class _MemoryObjectStore:
    def __init__(self, objects: dict[str, bytes]) -> None:
        self.objects = objects

    async def put(self, key: str, source: AsyncIterator[bytes]) -> StoredObject:
        content = b"".join([chunk async for chunk in source])
        self.objects[key] = content
        return StoredObject(key, len(content), sha256(content).hexdigest())

    async def open(self, key: str) -> AsyncIterator[bytes]:
        yield self.objects[key]

    async def delete(self, key: str) -> None:
        self.objects.pop(key, None)


def _profile(tmp_path: Path) -> OcrProfileSpec:
    directories = {role: tmp_path / role for role in ("detection", "recognition", "table")}
    for directory in directories.values():
        directory.mkdir()
    return OcrProfileSpec.create(
        pipeline_name="PP-StructureV3",
        pipeline_version="3.7.0",
        detection_model_name="PP-OCRv5_server_det",
        recognition_model_name="korean_PP-OCRv5_mobile_rec",
        table_model_name="SLANet_plus",
        languages=("ko", "en"),
        confidence_threshold=0.8,
        artifact_directories=directories,
    )


@pytest.mark.asyncio
async def test_docx_ocr_evidence_round_trips_to_authorized_image_viewer(
    tmp_path: Path,
) -> None:
    image_path = tmp_path / "policy.png"
    image_path.write_bytes(_PNG)
    source_path = tmp_path / "policy.docx"
    source = Document()
    source.add_heading("운용 기준", level=1)
    source.add_picture(str(image_path))
    source.save(source_path)

    asset_version_id = uuid4()
    projection_id = uuid4()
    parsed = DocxStructureParser(
        ocr_runtime=_FakeOcrRuntime(),
        ocr_profile=_profile(tmp_path),
    ).parse(ParseRequest(source_path, DOCX_MEDIA_TYPE, source_path.name, asset_version_id))
    chunked = StructuralChunker(_CharacterCounter()).chunk(
        parsed,
        projection_id=projection_id,
        config=ChunkingConfig(target_tokens=100, overlap_tokens=10, hard_ceiling_tokens=120),
    )

    evidence = next(
        item
        for item in chunked.evidence_units
        if item.location.source_kind is SourceKind.DOCX_IMAGE
    )
    assert evidence.location.source_kind is SourceKind.DOCX_IMAGE
    assert evidence.location.image_sha256 is not None

    original = source_path.read_bytes()
    normalized = serialize_parsed_document(parsed)
    actor_id = uuid4()
    resource = ViewerResource(
        document_id=uuid4(),
        asset_version_id=asset_version_id,
        asset_version_number=1,
        workspace_id=uuid4(),
        folder_id=None,
        projection_id=projection_id,
        title="합성 운용 기준",
        media_type=DOCX_MEDIA_TYPE,
        original_object_key="original",
        original_size=len(original),
        original_sha256=sha256(original).hexdigest(),
        parsed_object_key="parsed",
        parsed_sha256=sha256(normalized).hexdigest(),
    )
    service = ViewerService(
        _ViewerRepository(actor_id, resource),
        _MemoryObjectStore({"original": original, "parsed": normalized}),
    )

    image = await service.docx_image(
        actor_id=actor_id,
        asset_version_id=asset_version_id,
        projection_id=projection_id,
        element_id=evidence.location.element_id,
        image_sha256=evidence.location.image_sha256,
    )
    assert image.content == _PNG
    assert image.media_type == "image/png"

    with pytest.raises(AppError) as unauthorized:
        await service.docx_image(
            actor_id=uuid4(),
            asset_version_id=asset_version_id,
            projection_id=projection_id,
            element_id=evidence.location.element_id,
            image_sha256=evidence.location.image_sha256,
        )
    assert unauthorized.value.status_code == 404
